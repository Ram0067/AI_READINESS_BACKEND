from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.pagesizes import A4
from docx import Document
from io import BytesIO


def generate_pdf_report(assessment):
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4)
    styles = getSampleStyleSheet()
    story = []

    story.append(Paragraph("<b>AI Readiness Assessment Report</b>", styles["Title"]))
    story.append(Spacer(1, 12))

    # User Info
    story.append(Paragraph(f"<b>Name:</b> {assessment.person_name}", styles["Normal"]))
    story.append(Paragraph(f"<b>Company:</b> {assessment.company_name}", styles["Normal"]))
    story.append(Paragraph(f"<b>Email:</b> {assessment.email}", styles["Normal"]))
    story.append(Paragraph(f"<b>Designation:</b> {assessment.designation}", styles["Normal"]))
    story.append(Spacer(1, 12))

    # Score
    story.append(Paragraph(f"<b>Overall Score:</b> {assessment.overall_score}%", styles["Normal"]))
    story.append(Paragraph(f"<b>Category:</b> {assessment.category}", styles["Normal"]))
    story.append(Spacer(1, 12))

    # Feedback
    story.append(Paragraph("<b>Summary</b>", styles["Heading2"]))
    story.append(Paragraph(assessment.feedback_summary, styles["Normal"]))

    story.append(Spacer(1, 8))
    story.append(Paragraph("<b>Detailed Feedback</b>", styles["Heading2"]))
    story.append(Paragraph(assessment.feedback_category_detail, styles["Normal"]))

    story.append(Spacer(1, 12))
    story.append(Paragraph("<b>Recommended Actions</b>", styles["Heading2"]))
    for act in assessment.feedback_recommended_actions:
        story.append(Paragraph(f"- {act}", styles["Normal"]))

    # Questions & Answers
    story.append(Spacer(1, 12))
    story.append(Paragraph("<b>Questions & Responses</b>", styles["Heading2"]))

    for ans in assessment.answers.all():
        story.append(Paragraph(f"<b>{ans.question.text}</b>", styles["Normal"]))
        story.append(Paragraph(f"Answer: {ans.raw_value}", styles["Italic"]))
        story.append(Spacer(1, 6))

    doc.build(story)
    buffer.seek(0)
    return buffer


def generate_word_report(assessment):
    doc = Document()
    doc.add_heading("AI Readiness Assessment Report", level=1)

    doc.add_paragraph(f"Name: {assessment.person_name}")
    doc.add_paragraph(f"Company: {assessment.company_name}")
    doc.add_paragraph(f"Email: {assessment.email}")
    doc.add_paragraph(f"Designation: {assessment.designation}")

    doc.add_heading("Score", level=2)
    doc.add_paragraph(f"Overall Score: {assessment.overall_score}%")
    doc.add_paragraph(f"Category: {assessment.category}")

    doc.add_heading("Summary", level=2)
    doc.add_paragraph(assessment.feedback_summary)

    doc.add_heading("Detailed Feedback", level=2)
    doc.add_paragraph(assessment.feedback_category_detail)

    doc.add_heading("Recommended Actions", level=2)
    for act in assessment.feedback_recommended_actions:
        doc.add_paragraph(f"- {act}", style="List Bullet")

    doc.add_heading("Questions & Answers", level=2)
    for ans in assessment.answers.all():
        doc.add_paragraph(ans.question.text, style="List Number")
        doc.add_paragraph(f"Answer: {ans.raw_value}")

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
